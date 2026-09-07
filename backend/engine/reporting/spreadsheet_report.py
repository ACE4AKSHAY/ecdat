"""
M9: Spreadsheet (CSV and XLSX) Export Submodule
Generates flat asset lists in CSV or Excel XLSX format using openpyxl.
"""
import io
import csv
from typing import List, Dict, Any, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("asset_id", "Asset ID"),
    ("bom_ref", "BOM Ref"),
    ("type", "Type"),
    ("name", "Algorithm"),
    ("file_path", "File/Location"),
    ("library", "Library"),
    ("key_size", "Key Size"),
    ("business_criticality", "Business Criticality"),
    ("exposure", "Exposure"),
    ("risk_tier", "Risk Tier"),
    ("risk_score", "Risk Score"),
    ("urgency_ratio_r", "Urgency Ratio (r)"),
    ("shelf_life_x", "Shelf Life (X)"),
    ("migration_effort_y", "Migration Effort (Y)"),
    ("threat_timeline_z", "Threat Timeline (Z)"),
    ("quantum_vulnerable", "Quantum Vulnerable"),
    ("recommended_replacement", "Recommended Algorithm"),
    ("complexity", "Migration Complexity"),
    ("reference_standard", "Reference Standard"),
]


def export_csv(scan_id: str, assets_data: List[Dict[str, Any]]) -> Tuple[bytes, str, str]:
    """Generates standard CSV document."""
    output = io.StringIO()
    writer = csv.writer(output)

    # Write Header
    writer.writerow([label for _, label in COLUMNS])

    # Write Rows
    for a in assets_data:
        writer.writerow([
            a.get("id", ""),
            a.get("bom_ref", ""),
            a.get("type", a.get("primitive_category", "Algorithm")),
            a.get("name", ""),
            f"{a.get('file_path', '')}{(':' + str(a.get('line_number'))) if a.get('line_number') else ''}",
            a.get("library", ""),
            a.get("key_size", ""),
            a.get("business_criticality", ""),
            a.get("exposure", ""),
            a.get("risk_tier", ""),
            a.get("risk_score", ""),
            a.get("urgency_ratio_r", ""),
            a.get("shelf_life_x", ""),
            a.get("migration_effort_y", ""),
            a.get("threat_timeline_z", ""),
            "Yes" if a.get("quantum_vulnerable") else "No",
            a.get("recommended_replacement", ""),
            a.get("complexity", a.get("migration_complexity", "")),
            a.get("reference_standard", ""),
        ])

    content = output.getvalue().encode("utf-8")
    filename = f"ecdat-assets-{scan_id}.csv"
    return content, "text/csv", filename


def export_xlsx(scan_id: str, assets_data: List[Dict[str, Any]]) -> Tuple[bytes, str, str]:
    """Generates styled Excel XLSX spreadsheet using openpyxl."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cryptographic Assets"

    # Styling definitions matching ECDAT brand kit
    header_fill = PatternFill(start_color="263A73", end_color="263A73", fill_type="solid") # Cipher blue
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cell_font = Font(name="Arial", size=9)
    mono_font = Font(name="Consolas", size=9)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="E1E5E8"),
        right=Side(style="thin", color="E1E5E8"),
        top=Side(style="thin", color="E1E5E8"),
        bottom=Side(style="thin", color="E1E5E8"),
    )

    # Risk Tier Fill Colors
    risk_fills = {
        "Critical": PatternFill(start_color="FBEAE9", end_color="FBEAE9", fill_type="solid"),
        "High": PatternFill(start_color="FBEEDF", end_color="FBEEDF", fill_type="solid"),
        "Medium": PatternFill(start_color="F8F1D8", end_color="F8F1D8", fill_type="solid"),
        "Low": PatternFill(start_color="E4F2EB", end_color="E4F2EB", fill_type="solid"),
    }
    risk_fonts = {
        "Critical": Font(name="Arial", size=9, bold=True, color="B3261E"),
        "High": Font(name="Arial", size=9, bold=True, color="B5590F"),
        "Medium": Font(name="Arial", size=9, bold=True, color="93790E"),
        "Low": Font(name="Arial", size=9, bold=True, color="2E7D5B"),
    }

    # Write Header Row
    for col_idx, (_, label) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    ws.row_dimensions[1].height = 24

    # Write Data Rows
    for row_idx, a in enumerate(assets_data, 2):
        loc = f"{a.get('file_path', '')}{(':' + str(a.get('line_number'))) if a.get('line_number') else ''}"
        tier = a.get("risk_tier", "Low")

        row_values = [
            a.get("id", ""),
            a.get("bom_ref", ""),
            a.get("type", a.get("primitive_category", "Algorithm")),
            a.get("name", ""),
            loc,
            a.get("library", ""),
            a.get("key_size", ""),
            a.get("business_criticality", ""),
            a.get("exposure", ""),
            tier,
            a.get("risk_score", 0.0),
            a.get("urgency_ratio_r", 0.0),
            a.get("shelf_life_x", 0.0),
            a.get("migration_effort_y", 0.0),
            a.get("threat_timeline_z", 8.0),
            "Yes" if a.get("quantum_vulnerable") else "No",
            a.get("recommended_replacement", ""),
            a.get("complexity", a.get("migration_complexity", "")),
            a.get("reference_standard", ""),
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = cell_font
            cell.alignment = left_align
            cell.border = thin_border

            # Highlight Risk Tier
            if col_idx == 10 and tier in risk_fills:
                cell.fill = risk_fills[tier]
                cell.font = risk_fonts[tier]
                cell.alignment = center_align

            # Mono styling for Algorithm & BOM Ref
            if col_idx in [2, 4]:
                cell.font = mono_font

        ws.row_dimensions[row_idx].height = 20

    # Auto-fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    content = output.getvalue()
    filename = f"ecdat-assets-{scan_id}.xlsx"
    return content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename
